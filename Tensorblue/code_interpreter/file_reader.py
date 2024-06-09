import PyPDF2
import pdfplumber
import openpyxl
import pandas as pd
import docx
import openai
import sys
import io

# Functions to read different file types
def read_pdf(file_path):
    try:
        # Open the PDF file using pdfplumber
        with pdfplumber.open(file_path) as pdf:
            text = ''   # Initialize an empty string to store the text content
            # Iterate over each page in the PDF file
            for page in pdf.pages:     
                text += page.extract_text()  # Extract the text from the current page and append it to the text variable
        return text
    except Exception as e:
        return str(e)

# Function to read an Excel file (XLSX) and extract its data
def read_xlsx(file_path):
    # Load the Excel file using openpyxl
    workbook = openpyxl.load_workbook(file_path)
    # Select the active sheet
    sheet = workbook.active
    # Initialize an empty list to store the data
    data = []
    # Iterate over each row in the sheet
    for row in sheet.iter_rows(values_only=True):
        # Append the row data to the list
        data.append(list(row))
    # Return the extracted data
    return data


# Function to read a CSV file and extract its data
def read_csv(file_path):
    # Read the CSV file using pandas
    df = pd.read_csv(file_path)
    # Convert the DataFrame to a string and return it
    return df.to_string()

# Function to read a Word document (DOCX) and extract its text content
def read_docx(file_path):
    # Open the Word document using docx
    doc = docx.Document(file_path)
    # Initialize an empty string to store the text content
    text = ''
    # Iterate over each paragraph in the document
    for para in doc.paragraphs:
        # Append the paragraph text to the text variable
        text += para.text
    # Return the extracted text content
    return text

# file_path = r'C:\Users\Arsh5\Desktop\ANJALI\MCA CONTENT\Unit1.pdf'
# pdf_content = read_pdf(file_path)
# print(pdf_content)

# Function to generate Python code using OpenAI GPT-3.5 API
def generate_code(prompt):
    # Set the OpenAI API key
    openai.api_key = 'your-api-key'  # Replace with your actual OpenAI API key
    # Create a completion request using the prompt
    response = openai.Completion.create(
        model="text-davinci-003",
        prompt=prompt,
        max_tokens=150,
        n=1,
        stop=None,
        temperature=0.7,
    )
    # Return the generated code
    return response.choices[0].text.strip()

# Function to execute the generated Python code
def execute_code(code):
    # Save the current stdout
    old_stdout = sys.stdout
    # Create a new stdout
    new_stdout = io.StringIO()
    # Redirect stdout to the new stdout
    sys.stdout = new_stdout
    # Execute the code
    try:
        exec(code)
    except Exception as e:
        # Return the error message if an exception occurs
        return str(e)
    # Get the output from the new stdout
    output = new_stdout.getvalue()
    # Restore the original stdout
    sys.stdout = old_stdout
    # Return the output
    return output


# Function to handle file uploads and process user prompts
def handle_file_upload(file_path, file_type, user_prompt):
    # Read the file content based on the file type
    if file_type == 'pdf':
        content = read_pdf(file_path)
    elif file_type == 'xlsx':
        content = read_xlsx(file_path)
    elif file_type == 'csv':
        content = read_csv(file_path)
    elif file_type == 'docx':
        content = read_docx(file_path)
    else:
        # Return an error message for unsupported file types
        return "Unsupported file type"
    # Create a full prompt by combining the user prompt and file content
    full_prompt = f"{user_prompt}\n\nFile content:\n{content}"
    # Generate code using the full prompt
    code = generate_code(full_prompt)
    # Execute the generated code
    output = execute_code(code)
    # Return the output
    return output


# Main function to read a file based on its type
def main(file_path, file_type):
    # Read the file content based on the file type
    if file_type == 'pdf':
        return read_pdf(file_path)
    elif file_type == 'xlsx':
        return read_xlsx(file_path)
    elif file_type == 'csv':
        return read_csv(file_path)
    elif file_type == 'docx':
        return read_docx(file_path)
    else:
        # Return an error message for unsupported file types
        return "Unsupported file type"

# Entry point of the script
if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python script.py <file_path> <file_type>")
        sys.exit(1)
    file_path = sys.argv[1]
    file_type = sys.argv[2]
    result = main(file_path, file_type)
    print(result)
